from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from datetime import time, timedelta
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


class ForteError(ValueError):
    """An actionable input/model error, never a negative prediction."""


def key(value: str) -> str:
    return Path(unicodedata.normalize('NFKC', str(value)).strip()).stem.casefold().replace("х", "x")


def timestamp(value) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, timedelta):
        result = value.total_seconds()
    elif isinstance(value, (float, int)) and not isinstance(value, bool):
        # Plain numeric cells mean seconds; Excel time-formatted cells arrive as time.
        result = float(value)
    else:
        s = str(value).lower().strip().replace(",", ".")
        if re.fullmatch(r"\d+(?::\d{1,2}){1,2}(?:\.\d+)?", s):
            parts = [float(p) for p in s.split(":")]
            if any(p >= 60 for p in parts[1:]):
                raise ForteError(f"Некорректный timestamp: {value}")
            result = sum(p * 60**i for i, p in enumerate(reversed(parts)))
        elif re.fullmatch(r"\d+(?:\.\d+)?", s):
            result = float(s)
        else:
            match = re.fullmatch(
                r"(?:(\d+)\s*(?:ч|час(?:а|ов)?)\s*)?"
                r"(?:(\d+)\s*(?:мин(?:ута|уты|ут)?|м)\s*)?"
                r"(?:(\d+(?:\.\d+)?)\s*(?:сек(?:унда|унды|унд)?|с)\s*)?", s
            )
            if not match or not any(match.groups()):
                raise ForteError(f"Неизвестный формат timestamp: {value}")
            result = sum(float(v or 0) * m for v, m in zip(match.groups(), (3600, 60, 1)))
    if not math.isfinite(result) or result < 0:
        raise ForteError(f"Некорректный timestamp: {value}")
    return result


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_labels(labels: Path, videos: Path) -> list[dict]:
    if not labels.is_file():
        raise ForteError(f"Нет разметки: {labels}")
    if not videos.is_dir():
        raise ForteError(f"Нет каталога видео: {videos}")
    index = {}
    for p in videos.iterdir():
        if p.is_file() and p.suffix.lower() == ".mov":
            k = key(p.name)
            if k in index:
                raise ForteError(f"Неоднозначные имена X/Х: {p.name}, {index[k].name}")
            index[k] = p
    try:
        wb = load_workbook(labels, read_only=True, data_only=True)
    except (OSError, BadZipFile, ValueError) as exc:
        raise ForteError(f'Невозможно прочитать Excel: {exc}') from exc
    try:
        rows = iter(wb.active.iter_rows(values_only=True))
        header = [str(v).strip().casefold() for v in next(rows, ())]
        expected = ["наименование файла", "таймстемп", "нарушение / без нарушения"]
        if any(header.count(v) != 1 for v in expected):
            raise ForteError(f"Ожидаются уникальные колонки: {expected}")
        columns = [header.index(v) for v in expected]
        result, seen = [], set()
        for number, row in enumerate(rows, 2):
            if all(v is None for v in row):
                continue
            name, ts, label = [row[i] if i < len(row) else None for i in columns]
            k = key(name or "")
            if not name or k not in index:
                raise ForteError(f"Строка {number}: видео не найдено: {name}")
            if k in seen:
                raise ForteError(f"Строка {number}: повтор видео {name}; объедините события явно")
            seen.add(k)
            labels_map = {"нарушение": 1, "без нарушения": 0, "1": 1, "0": 0}
            lab = str(label).strip().casefold()
            if lab not in labels_map:
                raise ForteError(f"Строка {number}: неизвестная метка {label}")
            y, seconds = labels_map[lab], timestamp(ts)
            if y == 1 and seconds is None:
                raise ForteError(f"Строка {number}: для нарушения нужен timestamp")
            if y == 0 and seconds is not None:
                raise ForteError(f"Строка {number}: timestamp у отрицательного видео требует проверки")
            result.append(dict(id=k, file=index[k].name, label=y, timestamp=seconds))
        if not result:
            raise ForteError("Разметка пуста")
        return result
    finally:
        wb.close()


def write_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    temporary.replace(path)


def prepare(labels: Path, videos: Path, output: Path):
    from sklearn.model_selection import train_test_split

    if output.exists():
        raise ForteError("Manifest уже существует. Используйте его; новый split требует нового пути.")
    source_rows = read_labels(labels, videos)
    grouped = {}
    for row in source_rows:
        digest = sha256(videos/row['file'])
        if digest not in grouped:
            grouped[digest] = dict(row, sha256=digest, aliases=[], timestamps=[] if row['timestamp'] is None else [row['timestamp']])
        else:
            canonical = grouped[digest]
            if row['label'] != canonical['label']:
                raise ForteError(f"Одинаковое содержимое с конфликтующими метками: {canonical['file']}, {row['file']}")
            canonical['aliases'].append(row['file'])
            if row['timestamp'] is not None:
                canonical['timestamps'] = sorted(set(canonical['timestamps']+[row['timestamp']]))
    rows = list(grouped.values())
    try:
        train, holdout = train_test_split(rows, test_size=0.4, random_state=42, stratify=[r['label'] for r in rows])
        val, test = train_test_split(holdout, test_size=0.5, random_state=42, stratify=[r['label'] for r in holdout])
    except ValueError as exc:
        raise ForteError(f"Недостаточно видео каждого класса для трех выборок: {exc}") from exc
    for split, group in (("train", train), ("validation", val), ("test", test)):
        for row in group:
            row.update(split=split)
    manifest = dict(version=1, seed=42, source_rows=len(source_rows), unique_videos=len(rows), labels_sha256=sha256(labels), rows=sorted(rows, key=lambda r: r['id']))
    validate_manifest(manifest, videos)
    write_json(output, manifest)
    return {s: sum(r['split'] == s for r in rows) for s in ('train', 'validation', 'test')}


def validate_manifest(manifest, videos: Path):
    if manifest.get('version') != 1 or not manifest.get('rows'):
        raise ForteError("Неверный manifest")
    seen, hashes = set(), set()
    for row in manifest['rows']:
        name = row['file']
        if Path(name).name != name or '/' in name or '\\' in name:
            raise ForteError("Manifest должен содержать только имена файлов")
        if row['id'] != key(name) or row['id'] in seen:
            raise ForteError("Повтор или некорректный id в manifest")
        if row['split'] not in {'train', 'validation', 'test', 'excluded'}:
            raise ForteError("Неизвестная выборка")
        if row['label'] not in (0, 1):
            raise ForteError("Неверная бинарная метка")
        ts = timestamp(row['timestamp'])
        if (row['label'] == 1 and ts is None) or (row['label'] == 0 and ts is not None):
            raise ForteError("Метка и timestamp противоречат друг другу")
        events = row.get('timestamps', [] if ts is None else [ts])
        if (row['label'] == 1 and (not events or ts not in events)) or (row['label'] == 0 and events):
            raise ForteError('Некорректный список timestamps')
        for event in events:
            if timestamp(event) is None:
                raise ForteError('Пустой timestamp в списке событий')
        path = (videos / name).resolve()
        if not path.is_relative_to(videos.resolve()) or not path.is_file():
            raise ForteError(f"Видео недоступно: {name}")
        actual = sha256(path)
        if actual != row['sha256']:
            raise ForteError(f"Видео изменилось после split: {name}")
        if actual in hashes:
            raise ForteError("Обнаружены одинаковые видео: риск утечки между выборками")
        for alias in row.get('aliases', []):
            if Path(alias).name != alias or '/' in alias or '\\' in alias:
                raise ForteError('Некорректное имя alias')
            alias_path = (videos/alias).resolve()
            if not alias_path.is_relative_to(videos.resolve()) or not alias_path.is_file() or sha256(alias_path) != actual:
                raise ForteError(f'Alias отсутствует или отличается от canonical video: {alias}')
        seen.add(row['id'])
        hashes.add(actual)
    for split in ('train', 'validation', 'test'):
        if {r['label'] for r in manifest['rows'] if r['split'] == split} != {0, 1}:
            raise ForteError(f"В {split} необходимы оба класса")
    return manifest['rows']
