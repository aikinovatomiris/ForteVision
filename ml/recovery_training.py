"""Recover typed supervision from a supplied workbook; frozen test stays outside fitting."""
import argparse
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter

import joblib
import numpy as np
from openpyxl import load_workbook

from ml.dataset import ForteError, key, sha256, validate_manifest, write_json
from ml.detection import ResNetExtractor
from ml.preprocessing import SamplingConfig
from ml.sequence import fit_sequence
from ml.typed import CATALOG, fit


def read_descriptions(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        for sheet in workbook:
            for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if not cells or not isinstance(cells[0], str):
                    continue
                ident = key(cells[0] + ('' if '.' in cells[0] else '.mov'))
                if not ident.startswith('x-') or not ident[2:].isdigit():
                    continue
                raw = str(cells[1] or '').strip().casefold() if len(cells) > 1 else ''
                if raw not in ('', 'есть', 'нет'):
                    raise ForteError(f'Unknown workbook label: {sheet.title}:{number}')
                row = dict(id=ident, binary_label={'есть': 1, 'нет': 0}.get(raw),
                           description=str(cells[2] or '').strip() if len(cells) > 2 else '',
                           source_sheet=sheet.title, source_row=number)
                if ident in result:
                    raise ForteError(f'Duplicate workbook ID: {ident}')
                result[ident] = row
    finally:
        workbook.close()
    return result


def reconcile(manifest, annotations, mapping):
    # Conservative groups from prior project review. Never move a test clip to train.
    protected_pairs = [('x-16', 'x-17'), ('x-22', 'x-23'), ('x-12', 'x-24')]
    test_ids = {r['id'] for r in manifest['rows'] if r['split'] == 'test'}
    val_ids = {r['id'] for r in manifest['rows'] if r['split'] == 'validation'}
    quarantine = set()
    for pair in protected_pairs:
        if set(pair) & test_ids:
            quarantine.update(set(pair)-test_ids)
        elif set(pair) & val_ids:
            quarantine.update(set(pair)-val_ids)
    output, conflicts = [], []
    for row in manifest['rows']:
        ids = [row['id']] + [key(name) for name in row.get('aliases', [])]
        descriptions = [annotations[i] for i in ids if i in annotations]
        labels = {r['binary_label'] for r in descriptions if r['binary_label'] is not None}
        if len(labels) > 1:
            raise ForteError(f'Conflicting workbook labels for identical video: {ids}')
        label = next(iter(labels)) if labels else None
        types = {mapping['types'][i] for i in ids if mapping['types'].get(i)}
        if len(types) > 1:
            raise ForteError(f'Conflicting category for identical video: {ids}')
        category = 'normal' if label == 0 else next(iter(types)) if label == 1 and types else None
        if category and category != 'normal' and category not in CATALOG:
            raise ForteError(f'Unknown reviewed taxonomy: {category}')
        if label is not None and label != row['label']:
            conflicts.append(dict(id=row['id'], original=row['label'], workbook=label,
                                  resolution='Latest user-supplied workbook; original preserved'))
        # Original times only when the binary sources agree. They are weak anchors.
        timestamps = row['timestamps'] if label == row['label'] == 1 else []
        override = mapping.get('timestamp_overrides', {}).get(row['id'])
        if override and label == 1:
            timestamps = override['timestamps']
        output.append(dict(row=row, label=category, binary_label=label, timestamps=timestamps,
                           timestamp_source=override['source'] if override else 'original_weak_anchor' if timestamps else 'unknown',
                           annotations=descriptions, excluded_from_fitting=row['id'] in quarantine,
                           reason='Related recording belongs to held-out split' if row['id'] in quarantine else
                                  'Missing or ambiguous event type' if category is None else None))
    return output, conflicts


def extract(records, root, directory, extractor, sampling):
    result = []
    for record in records:
        row = record['row']
        started = perf_counter()
        print(f"extract {row['split']} {row['id']}", flush=True)
        cache = directory/'features'/f"{row['sha256']}.joblib"
        identity = dict(video_sha256=row['sha256'], weights_sha256=extractor.weights_sha256,
                        sampling=sampling.to_dict(), image_mode='full_frame')
        if cache.exists():
            data = joblib.load(cache)
            if data['identity'] != identity:
                raise ForteError('Feature cache identity mismatch')
        else:
            x, windows, info = extractor.extract(root/'data/videos'/row['file'], sampling, started+285)
            data = dict(identity=identity, x=x, windows=windows, info=info)
            cache.parent.mkdir(parents=True, exist_ok=True)
            joblib.dump(data, cache)
        if any(t < 0 or t >= data['info']['duration_seconds'] for t in record['timestamps']):
            raise ForteError(f"Timestamp outside video: {row['id']}")
        eligible = np.ones(len(data['x']), dtype=bool)
        if record['binary_label'] == 1 and record['timestamps']:
            eligible = np.array([any(w['start'] <= t+1 and w['end'] > t-1 for t in record['timestamps'])
                                 for w in data['windows']])
        result.append(dict(record, x=data['x'], windows=data['windows'], eligible=eligible,
                           weak=record['binary_label'] == 1 and not record['timestamps'],
                           elapsed_seconds=perf_counter()-started))
    return result


def evaluate(model, records, threshold):
    counts = dict(tp=0, fp=0, tn=0, fn=0)
    details, errors = [], []
    typed_count = typed_hits = 0
    for record in records:
        events = model.events(record['x'], record['windows'], threshold)
        label = record['binary_label']
        if label is not None:
            counts[('tp' if events else 'fn') if label else ('fp' if events else 'tn')] += 1
        expected = record['label']
        correct = expected is not None and expected != 'normal' and any(e.get('event_type') == expected for e in events)
        if label == 1 and expected:
            typed_count += 1
            typed_hits += correct
        distances = [min(abs(e['peak']-t) for e in events) for t in record['timestamps']] if events else []
        errors.extend(distances)
        details.append(dict(id=record['row']['id'], expected_type=expected, binary_label=label,
            expected_type_seen=expected in model.estimator.classes_, correct_type=bool(correct),
            annotation_times=record['timestamps'], timestamp_source=record['timestamp_source'],
            events=[dict(type=e.get('event_type'), time=e['peak'], start=e['start'], end=e['end'], score=e['score']) for e in events]))
    tp, fp, fn, tn = (counts[k] for k in ('tp', 'fp', 'fn', 'tn'))
    return dict(**counts, confusion_matrix=[[tn, fp], [fn, tp]],
        precision=tp/(tp+fp) if tp+fp else None, recall=tp/(tp+fn) if tp+fn else None,
        f1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else 0.,
        correctly_typed_positive_videos=int(typed_hits), typed_positive_videos=typed_count,
        conditional_nearest_peak_mae=float(np.mean(errors)) if errors else None,
        localization_note='Weak original anchors except explicit overrides; misses excluded from conditional MAE.',
        videos=details)


def run(root, workbook, mapping_path, directory, target, weights, test_only=False):
    manifest_path = root/'artifacts/manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    validate_manifest(manifest, root/'data/videos')
    annotations = read_descriptions(workbook)
    mapping = json.loads(mapping_path.read_text(encoding='utf-8'))
    records, conflicts = reconcile(manifest, annotations, mapping)
    provenance = dict(workbook_sha256=sha256(workbook), mapping_sha256=sha256(mapping_path),
                      manifest_sha256=sha256(manifest_path))
    if not test_only and (target.exists() or target.with_suffix('.json').exists()):
        raise ForteError('Refusing to overwrite model; use a fresh output path')
    if not test_only:
        write_json(directory/'annotation_audit.json', dict(provenance=provenance, conflicts=conflicts,
            records=records, unmatched_workbook_ids=sorted(set(annotations)-{i for r in manifest['rows'] for i in [r['id']]+[key(a) for a in r.get('aliases', [])]})))
    sampling = SamplingConfig(fps=2, window_seconds=4, stride_seconds=2, batch_size=16)
    extractor = ResNetExtractor(weights, 'cpu', 'full_frame')
    extractor.torch.manual_seed(42)
    np.random.seed(42)
    if test_only:
        metadata = json.loads(target.with_suffix('.json').read_text(encoding='utf-8'))
        if any(metadata.get(k) != v for k, v in provenance.items()) or sha256(target) != metadata['pipeline_sha256']:
            raise ForteError('Frozen model/provenance changed before evaluation')
        test = [r for r in records if r['row']['split'] == 'test']
        if set(metadata['train_sha256']) & {r['row']['sha256'] for r in test}:
            raise ForteError('Test leakage')
        report = evaluate(joblib.load(target), extract(test, root, directory, extractor, sampling), metadata['threshold'])
        write_json(directory/f'{target.stem}_test_evaluation.json', dict(model_sha256=metadata['pipeline_sha256'],
            report=report, note='Frozen evaluation; previously viewed internal test, not independent new-camera data.'))
        print(json.dumps(report, ensure_ascii=True, indent=2))
        return
    train_rows = [r for r in records if r['row']['split'] == 'train' and r['label'] is not None and not r['excluded_from_fitting']]
    val_rows = [r for r in records if r['row']['split'] == 'validation' and r['binary_label'] is not None]
    if 'normal' not in {r['label'] for r in train_rows} or {r['binary_label'] for r in val_rows} != {0, 1}:
        raise ForteError('Training normal class and both validation labels are required')
    training = extract(train_rows, root, directory, extractor, sampling)
    validation = extract(val_rows, root, directory, extractor, sampling)
    candidates = []
    for name, fitter in [('appearance', fit), ('sequence', fit_sequence)]:
        print('fit', name, flush=True)
        model = fitter(training)
        # Small, predeclared search on validation only, not on X-4 or test.
        for threshold in (.35, .55, .75):
            report = evaluate(model, validation, threshold)
            candidates.append((model, name, threshold, report))
    chosen = max(candidates, key=lambda item: (item[3]['f1'], -item[3]['fp'], item[2], item[1] == 'appearance'))
    model, name, threshold, report = chosen
    target.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(model, target)
    metadata = dict(schema_version=1, architecture='resnet18', pooling='mean_max', feature_dim=1024,
        embedding_dim=512, weights_name='IMAGENET1K_V1', weights_sha256=extractor.weights_sha256,
        pipeline_sha256=sha256(target), image_mode='full_frame', sampling=sampling.to_dict(),
        labeling=dict(positive_radius=1, exclusion_zone=6, weak_negatives=False),
        threshold=threshold, threshold_objective='validation_video_f1_then_fewer_false_positives',
        classifier=name, created_at=datetime.now(timezone.utc).isoformat(), seed=42, **provenance,
        train_ids=[r['row']['id'] for r in training], train_sha256=[r['row']['sha256'] for r in training],
        validation_ids=[r['row']['id'] for r in validation], test_used_for_training=False,
        event_types=[str(c) for c in model.estimator.classes_ if c != 'normal'],
        training_video_counts=dict(Counter(r['label'] for r in training)),
        review_only=True, validation_metrics=report,
        limitations=['Few recordings per type; descriptions are weak supervision, not framewise action labels.',
                    'Class names/descriptions are hypotheses, not verified object detections.',
                    'No independent test or new-camera accuracy claim; no automatic legal finding.'])
    write_json(target.with_suffix('.json'), metadata)
    write_json(directory/'selection.json', dict(selected=dict(architecture=name, threshold=threshold, report=report),
        candidates=[dict(architecture=n, threshold=t, report=r) for _, n, t, r in candidates],
        note='Selection on a small reused validation set is optimistic. Test not extracted during training.'))
    regression = evaluate(model, [r for r in training if r['row']['id'] == 'x-4'], threshold)
    write_json(directory/'x4_training_regression.json', dict(report=regression, note='Known training example; not an accuracy test.'))
    print(json.dumps(dict(model=str(target), selected=name, threshold=threshold,
        train_counts=metadata['training_video_counts'], validation=report,
        x4_training_regression=regression), ensure_ascii=True, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--mapping', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, default=Path('artifacts/recovery_20261018'))
    parser.add_argument('--model', type=Path, default=Path('models/recovery_typed_v1.joblib'))
    parser.add_argument('--weights', type=Path, default=Path('models/resnet18-f37072fd.pth'))
    parser.add_argument('--evaluate-test', action='store_true')
    args = parser.parse_args()
    run(Path(__file__).resolve().parents[1], args.workbook, args.mapping, args.output_dir, args.model, args.weights, args.evaluate_test)


if __name__ == '__main__':
    main()
